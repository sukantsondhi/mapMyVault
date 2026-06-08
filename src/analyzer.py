"""Local content extraction and deterministic relationship evidence."""

from pathlib import Path
from typing import Dict, Iterable, List, Optional, Set
import ast
import csv
import hashlib
import io
import json
import re
import shutil
import os

import pathspec


TEXT_EXTENSIONS = {
    ".py", ".js", ".jsx", ".ts", ".tsx", ".java", ".cs", ".go", ".rs", ".rb",
    ".php", ".c", ".h", ".cpp", ".hpp", ".md", ".txt", ".rst", ".json", ".yaml",
    ".yml", ".toml", ".ini", ".cfg", ".xml", ".html", ".css", ".scss", ".sql",
    ".sh", ".ps1", ".bat", ".dockerfile",
}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx", ".csv", ".tsv"}
OCR_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}
PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPO_MODELS_DIR = PROJECT_ROOT / "models"


def stable_id(path: str) -> str:
    return hashlib.sha256(path.encode("utf-8")).hexdigest()[:24]


def content_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def is_probably_binary(data: bytes) -> bool:
    return b"\x00" in data[:4096]


def load_ignore_spec(root: Path, extra_patterns: Iterable[str]):
    patterns = list(extra_patterns)
    gitignore = root / ".gitignore"
    if gitignore.exists():
        patterns.extend(gitignore.read_text(encoding="utf-8", errors="ignore").splitlines())
    return pathspec.GitIgnoreSpec.from_lines(patterns)


def scan_repository(root: Path, output: Path, excludes: Iterable[str], max_bytes: int):
    output_resolved = output.resolve()
    spec = load_ignore_spec(root, excludes)
    records = []
    for item in sorted(root.rglob("*")):
        resolved = item.resolve()
        if resolved == output_resolved or output_resolved in resolved.parents:
            continue
        rel = item.relative_to(root).as_posix()
        if spec.match_file(rel + ("/" if item.is_dir() else "")):
            continue
        stat = item.stat()
        record = {
            "id": stable_id(rel),
            "path": rel,
            "parent": item.parent.relative_to(root).as_posix()
            if item.parent != root
            else None,
            "kind": "folder" if item.is_dir() else "file",
            "extension": item.suffix.lower(),
            "size": 0 if item.is_dir() else stat.st_size,
            "content_hash": None,
            "modified_ns": stat.st_mtime_ns,
        }
        if item.is_file() and stat.st_size <= max_bytes:
            data = item.read_bytes()
            record["content_hash"] = content_hash(data)
        records.append(record)
    return records


def _trim(text: str, max_chars: int) -> str:
    return text.strip()[:max_chars]


def _finalize_extraction(text: str, properties: Dict, max_chars: int) -> tuple[str, Dict]:
    clean = _trim(text, max_chars)
    properties["extracted_characters"] = len(clean)
    properties["truncated"] = len(text.strip()) > max_chars
    return clean, properties


def _clean_ocr_language(language: str) -> str:
    return (language or "eng").strip() or "eng"


def _image_properties(path: Path) -> Dict:
    properties = {"format": path.suffix.lower().lstrip(".")}
    try:
        from PIL import Image

        with Image.open(path) as image:
            properties.update(
                {
                    "image_width": image.width,
                    "image_height": image.height,
                    "image_mode": image.mode,
                }
            )
    except Exception as exc:
        properties["image_metadata_error"] = str(exc)
    return properties


def _resolve_local_yolo_model(model: str) -> Optional[Path]:
    candidates = []
    env_model = os.environ.get("MAPMYVAULT_YOLO_MODEL")
    if env_model:
        candidates.append(Path(env_model).expanduser())
    if model:
        model_path = Path(model).expanduser()
        candidates.extend(
            [
                REPO_MODELS_DIR / model,
                model_path,
                Path.cwd() / model,
                Path.home() / ".cache" / "mapmyvault" / "models" / model,
                Path.home() / ".cache" / "ultralytics" / model,
            ]
        )
    for candidate in candidates:
        if candidate.is_file():
            return candidate.resolve()
    return None


def _analyze_image_yolo(
    path: Path,
    model_name: str,
    confidence: float,
    max_detections: int,
) -> tuple[str, Dict]:
    if not model_name:
        return "", {
            "vision_engine": "yolo",
            "vision_status": "vision_unavailable",
            "vision_error": "No local YOLO model path was configured.",
        }
    model_path = _resolve_local_yolo_model(model_name)
    if model_path is None:
        return "", {
            "vision_engine": "yolo",
            "vision_model": model_name,
            "vision_status": "vision_unavailable",
            "vision_error": (
                "YOLO weights were not found locally. Put the .pt file in the "
                "repo models folder, set MAPMYVAULT_YOLO_MODEL, or configure an "
                "absolute model path. "
                "mapMyVault will not auto-download model weights during indexing."
            ),
        }
    try:
        from ultralytics import YOLO
    except ImportError as exc:
        return "", {
            "vision_engine": "yolo",
            "vision_model": str(model_path),
            "vision_status": "vision_unavailable",
            "vision_error": f"Missing Python vision dependency: {exc.name}",
        }

    try:
        model = YOLO(str(model_path))
        results = model.predict(
            source=str(path),
            conf=float(confidence),
            max_det=int(max_detections),
            verbose=False,
        )
        detections = []
        label_counts: Dict[str, int] = {}
        for result in results:
            names = result.names or getattr(model, "names", {}) or {}
            for box in getattr(result, "boxes", []) or []:
                class_id = int(box.cls[0])
                label = str(names.get(class_id, class_id))
                score = float(box.conf[0])
                xyxy = [round(float(value), 2) for value in box.xyxy[0].tolist()]
                detections.append(
                    {
                        "label": label,
                        "confidence": round(score, 4),
                        "bbox_xyxy": xyxy,
                    }
                )
                label_counts[label] = label_counts.get(label, 0) + 1

        labels = sorted(label_counts)
        status = "vision_extracted" if detections else "vision_empty"
        text = ""
        if detections:
            label_text = ", ".join(
                f"{label} x{count}" for label, count in sorted(label_counts.items())
            )
            evidence = "; ".join(
                f"{item['label']} ({item['confidence']:.2f})"
                for item in detections[:10]
            )
            text = (
                "YOLO visual analysis:\n"
                f"Detected objects: {label_text}.\n"
                f"Top detections: {evidence}."
            )
        else:
            text = (
                "YOLO visual analysis:\n"
                "No objects from the configured YOLO model classes were detected. "
                "This does not mean the image is blank; it means the local YOLO "
                "model did not match any trained object class above the configured "
                "confidence threshold."
            )
        return text, {
            "vision_engine": "yolo",
            "vision_model": str(model_path),
            "vision_status": status,
            "vision_confidence": float(confidence),
            "vision_max_detections": int(max_detections),
            "vision_labels": labels,
            "vision_label_counts": label_counts,
            "vision_detections": detections[: int(max_detections)],
        }
    except Exception as exc:
        return "", {
            "vision_engine": "yolo",
            "vision_model": str(model_path),
            "vision_status": "vision_failed",
            "vision_error": str(exc),
        }


def _ocr_image_tesseract(path: Path, max_chars: int, language: str) -> tuple[str, Dict]:
    language = _clean_ocr_language(language)
    try:
        from PIL import Image
        import pytesseract
    except ImportError as exc:
        return "", {
            "format": path.suffix.lower().lstrip("."),
            "extraction_status": "ocr_unavailable",
            "ocr_engine": "tesseract",
            "ocr_error": f"Missing Python OCR dependency: {exc.name}",
            "extracted_characters": 0,
            "truncated": False,
        }
    if shutil.which("tesseract") is None:
        return "", {
            "format": path.suffix.lower().lstrip("."),
            "extraction_status": "ocr_unavailable",
            "ocr_engine": "tesseract",
            "ocr_language": language,
            "ocr_error": "Tesseract executable was not found in PATH.",
            "extracted_characters": 0,
            "truncated": False,
        }
    try:
        with Image.open(path) as image:
            text = pytesseract.image_to_string(image, lang=language)
        return _finalize_extraction(text, {
            "format": path.suffix.lower().lstrip("."),
            "extraction_status": "ocr_extracted" if text.strip() else "ocr_empty",
            "ocr_engine": "tesseract",
            "ocr_language": language,
        }, max_chars)
    except Exception as exc:
        return "", {
            "format": path.suffix.lower().lstrip("."),
            "extraction_status": "ocr_failed",
            "ocr_engine": "tesseract",
            "ocr_language": language,
            "ocr_error": str(exc),
            "extracted_characters": 0,
            "truncated": False,
        }


def _extract_image(
    path: Path,
    max_chars: int,
    enable_ocr: bool,
    ocr_engine: str,
    ocr_language: str,
    enable_vision: bool,
    vision_engine: str,
    vision_model: str,
    vision_confidence: float,
    vision_max_detections: int,
) -> tuple[str, Dict]:
    metadata = _image_properties(path)
    parts = []

    if enable_ocr:
        if ocr_engine != "tesseract":
            metadata.update(
                {
                    "extraction_status": "ocr_failed",
                    "ocr_status": "ocr_failed",
                    "ocr_error": f"Unsupported OCR engine: {ocr_engine}",
                    "extracted_characters": 0,
                    "truncated": False,
                }
            )
        else:
            ocr_text, ocr_metadata = _ocr_image_tesseract(
                path, max_chars, ocr_language
            )
            metadata.update(ocr_metadata)
            metadata["ocr_status"] = ocr_metadata.get("extraction_status")
            if ocr_text:
                parts.append(
                    f"OCR text:\n{ocr_text}" if enable_vision else ocr_text
                )

    if enable_vision:
        if vision_engine != "yolo":
            metadata.update(
                {
                    "vision_engine": vision_engine,
                    "vision_status": "vision_failed",
                    "vision_error": f"Unsupported vision engine: {vision_engine}",
                }
            )
        else:
            vision_text, vision_metadata = _analyze_image_yolo(
                path,
                vision_model,
                vision_confidence,
                vision_max_detections,
            )
            metadata.update(vision_metadata)
            if vision_text:
                parts.append(vision_text)

    if parts:
        ocr_status = metadata.get("ocr_status")
        vision_status = metadata.get("vision_status")
        if ocr_status in {"ocr_failed", "ocr_unavailable"}:
            metadata["extraction_status"] = ocr_status
        elif ocr_status == "ocr_extracted" and vision_status == "vision_extracted":
            metadata["extraction_status"] = "image_extracted"
        elif ocr_status == "ocr_extracted":
            metadata["extraction_status"] = "ocr_extracted"
        elif vision_status == "vision_extracted":
            metadata["extraction_status"] = "vision_extracted"
        else:
            metadata["extraction_status"] = "image_extracted"
        return _finalize_extraction("\n\n".join(parts), metadata, max_chars)

    if not metadata.get("extraction_status"):
        if enable_vision:
            metadata["extraction_status"] = metadata.get(
                "vision_status", "vision_failed"
            )
        elif enable_ocr:
            metadata["extraction_status"] = metadata.get(
                "ocr_status", "ocr_failed"
            )
        else:
            metadata["extraction_status"] = "unsupported_binary"
    metadata["extracted_characters"] = 0
    metadata["truncated"] = False
    return "", metadata


def _pdf_images_with_pdf2image(path: Path, dpi: int, max_pages: int):
    from pdf2image import convert_from_path

    return convert_from_path(
        str(path),
        dpi=dpi,
        first_page=1,
        last_page=max_pages if max_pages > 0 else None,
    )


def _pdf_images_with_pymupdf(path: Path, dpi: int, max_pages: int):
    try:
        import fitz
        from PIL import Image
    except ImportError as exc:
        raise RuntimeError(f"Missing PDF OCR renderer dependency: {exc.name}") from exc

    scale = dpi / 72
    images = []
    document = fitz.open(str(path))
    try:
        page_limit = len(document) if max_pages <= 0 else min(max_pages, len(document))
        for page_number in range(page_limit):
            page = document.load_page(page_number)
            pixmap = page.get_pixmap(
                matrix=fitz.Matrix(scale, scale),
                alpha=False,
            )
            images.append(
                Image.frombytes(
                    "RGB",
                    [pixmap.width, pixmap.height],
                    pixmap.samples,
                )
            )
    finally:
        document.close()
    return images


def _pdf_images_for_ocr(path: Path, dpi: int, max_pages: int):
    errors = []
    try:
        return _pdf_images_with_pdf2image(path, dpi, max_pages), "pdf2image"
    except Exception as exc:
        errors.append(f"pdf2image/poppler failed: {exc}")
    try:
        return _pdf_images_with_pymupdf(path, dpi, max_pages), "pymupdf"
    except Exception as exc:
        errors.append(f"pymupdf failed: {exc}")
    raise RuntimeError("; ".join(errors))


def _ocr_pdf_tesseract(
    path: Path,
    max_chars: int,
    language: str,
    dpi: int,
    max_pages: int,
) -> tuple[str, Dict]:
    language = _clean_ocr_language(language)
    try:
        import pytesseract
    except ImportError as exc:
        return "", {
            "format": "pdf",
            "extraction_status": "ocr_unavailable",
            "ocr_engine": "tesseract",
            "ocr_error": f"Missing Python OCR dependency: {exc.name}",
            "extracted_characters": 0,
            "truncated": False,
        }
    if shutil.which("tesseract") is None:
        return "", {
            "format": "pdf",
            "extraction_status": "ocr_unavailable",
            "ocr_engine": "tesseract",
            "ocr_language": language,
            "ocr_error": "Tesseract executable was not found in PATH.",
            "extracted_characters": 0,
            "truncated": False,
        }
    try:
        images, renderer = _pdf_images_for_ocr(path, dpi, max_pages)
        parts = []
        for image in images:
            parts.append(pytesseract.image_to_string(image, lang=language))
            if sum(len(item) for item in parts) >= max_chars:
                break
        text = "\n\n".join(parts)
        return _finalize_extraction(text, {
            "format": "pdf",
            "extraction_status": "ocr_extracted" if text.strip() else "ocr_empty",
            "ocr_engine": "tesseract",
            "ocr_language": language,
            "ocr_dpi": dpi,
            "ocr_pages": len(images),
            "ocr_max_pages": max_pages,
            "ocr_pdf_renderer": renderer,
        }, max_chars)
    except Exception as exc:
        return "", {
            "format": "pdf",
            "extraction_status": "ocr_failed",
            "ocr_engine": "tesseract",
            "ocr_language": language,
            "ocr_error": str(exc),
            "extracted_characters": 0,
            "truncated": False,
        }


def _extract_pdf(
    path: Path,
    max_chars: int,
    enable_ocr: bool = False,
    ocr_language: str = "eng",
    ocr_dpi: int = 200,
    ocr_max_pages: int = 10,
) -> tuple[str, Dict]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
        if sum(len(item) for item in pages) >= max_chars:
            break
    properties = {
        "format": "pdf",
        "page_count": len(reader.pages),
        "title": str((reader.metadata or {}).get("/Title") or ""),
        "author": str((reader.metadata or {}).get("/Author") or ""),
        "subject": str((reader.metadata or {}).get("/Subject") or ""),
    }
    text = "\n\n".join(pages)
    if text.strip():
        properties["extraction_status"] = "extracted"
        return _finalize_extraction(text, properties, max_chars)
    if enable_ocr:
        ocr_text, ocr_properties = _ocr_pdf_tesseract(
            path, max_chars, ocr_language, ocr_dpi, ocr_max_pages
        )
        ocr_properties.update({
            "page_count": len(reader.pages),
            "title": properties["title"],
            "author": properties["author"],
            "subject": properties["subject"],
        })
        return ocr_text, ocr_properties
    properties["extraction_status"] = "ocr_required"
    return _finalize_extraction(text, properties, max_chars)


def _extract_docx(path: Path, max_chars: int) -> tuple[str, Dict]:
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    core = document.core_properties
    return _finalize_extraction("\n".join(parts), {
        "format": "docx",
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "title": core.title or "",
        "author": core.author or "",
        "subject": core.subject or "",
        "keywords": core.keywords or "",
        "extraction_status": "extracted",
    }, max_chars)


def _extract_xlsx(path: Path, max_chars: int) -> tuple[str, Dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        parts = []
        sheet_rows = {}
        for sheet in workbook.worksheets:
            parts.append(f"Sheet: {sheet.title}")
            count = 0
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    parts.append(" | ".join(values))
                    count += 1
                if sum(len(item) for item in parts) >= max_chars:
                    break
            sheet_rows[sheet.title] = count
        return _finalize_extraction("\n".join(parts), {
            "format": "xlsx",
            "sheet_names": list(workbook.sheetnames),
            "rows_read_by_sheet": sheet_rows,
            "extraction_status": "extracted",
        }, max_chars)
    finally:
        workbook.close()


def _extract_pptx(path: Path, max_chars: int) -> tuple[str, Dict]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts = []
    for number, slide in enumerate(presentation.slides, 1):
        parts.append(f"Slide {number}")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
        if sum(len(item) for item in parts) >= max_chars:
            break
    return _finalize_extraction("\n".join(parts), {
        "format": "pptx",
        "slide_count": len(presentation.slides),
        "extraction_status": "extracted",
    }, max_chars)


def _extract_delimited(path: Path, max_chars: int) -> tuple[str, Dict]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    rows = []
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    for row in reader:
        rows.append(" | ".join(row))
        if sum(len(item) for item in rows) >= max_chars:
            break
    return _finalize_extraction("\n".join(rows), {
        "format": path.suffix.lower().lstrip("."),
        "rows_read": len(rows),
        "extraction_status": "extracted",
    }, max_chars)


def extract_content(
    path: Path,
    max_bytes: int,
    max_chars: int,
    enable_ocr: bool = False,
    ocr_engine: str = "tesseract",
    ocr_language: str = "eng",
    ocr_dpi: int = 200,
    ocr_max_pages: int = 10,
    enable_vision: bool = False,
    vision_engine: str = "yolo",
    vision_model: str = "yolov8n.pt",
    vision_confidence: float = 0.25,
    vision_max_detections: int = 50,
) -> tuple[str, Dict]:
    if path.stat().st_size > max_bytes:
        return "", {
            "format": path.suffix.lower().lstrip("."),
            "extraction_status": "too_large",
            "extracted_characters": 0,
            "truncated": False,
        }
    ocr_language = _clean_ocr_language(ocr_language)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        if enable_ocr and ocr_engine != "tesseract":
            return "", {
                "format": "pdf",
                "extraction_status": "ocr_failed",
                "ocr_error": f"Unsupported OCR engine: {ocr_engine}",
                "extracted_characters": 0,
                "truncated": False,
            }
        return _extract_pdf(
            path,
            max_chars,
            enable_ocr,
            ocr_language,
            ocr_dpi,
            ocr_max_pages,
        )
    if suffix == ".docx":
        return _extract_docx(path, max_chars)
    if suffix == ".xlsx":
        return _extract_xlsx(path, max_chars)
    if suffix == ".pptx":
        return _extract_pptx(path, max_chars)
    if suffix in {".csv", ".tsv"}:
        return _extract_delimited(path, max_chars)
    if (enable_ocr or enable_vision) and suffix in OCR_IMAGE_EXTENSIONS:
        return _extract_image(
            path,
            max_chars,
            enable_ocr,
            ocr_engine,
            ocr_language,
            enable_vision,
            vision_engine,
            vision_model,
            vision_confidence,
            vision_max_detections,
        )
    data = path.read_bytes()
    if is_probably_binary(data):
        if (enable_ocr or enable_vision) and suffix in OCR_IMAGE_EXTENSIONS:
            return _extract_image(
                path,
                max_chars,
                enable_ocr,
                ocr_engine,
                ocr_language,
                enable_vision,
                vision_engine,
                vision_model,
                vision_confidence,
                vision_max_detections,
            )
        return "", {"format": suffix.lstrip("."), "extraction_status": "unsupported_binary", "extracted_characters": 0, "truncated": False}
    if path.suffix.lower() not in TEXT_EXTENSIONS and path.name.lower() != "dockerfile":
        return "", {"format": suffix.lstrip("."), "extraction_status": "unsupported", "extracted_characters": 0, "truncated": False}
    text = data.decode("utf-8-sig", errors="replace")
    return _finalize_extraction(text, {
        "format": suffix.lstrip(".") or path.name.lower(),
        "line_count": text.count("\n") + 1,
        "extraction_status": "extracted",
    }, max_chars)


def extract_text(path: Path, max_bytes: int, max_chars: int) -> str:
    """Compatibility wrapper returning only locally extracted text."""
    return extract_content(path, max_bytes, max_chars)[0]


def parse_content(path: str, text: str) -> Dict[str, List[str]]:
    result = {"imports": [], "symbols": [], "references": []}
    suffix = Path(path).suffix.lower()
    if suffix == ".py":
        try:
            tree = ast.parse(text)
            for node in ast.walk(tree):
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                    result["symbols"].append(node.name)
                elif isinstance(node, ast.Import):
                    result["imports"].extend(alias.name for alias in node.names)
                elif isinstance(node, ast.ImportFrom) and node.module:
                    result["imports"].append(node.module)
        except SyntaxError:
            pass
    else:
        result["imports"].extend(
            re.findall(
                r"(?:from|import|require\()\s*[\"']?([A-Za-z0-9_./@-]+)",
                text,
            )
        )
        result["symbols"].extend(
            re.findall(r"\b(?:class|def|function|interface|struct)\s+([A-Za-z_]\w*)", text)
        )
    result["references"] = re.findall(
        r"(?:\[[^\]]+\]\(([^)]+)\)|[\"']([^\"']+\.[A-Za-z0-9]{1,8})[\"'])", text
    )
    result["references"] = [a or b for a, b in result["references"]]
    return {key: sorted(set(values))[:100] for key, values in result.items()}


def deterministic_edges(files: List[Dict]) -> List[Dict]:
    by_path = {row["path"]: row for row in files}
    by_stem: Dict[str, Set[str]] = {}
    for row in files:
        by_stem.setdefault(Path(row["path"]).stem.lower(), set()).add(row["id"])
    edges = []
    for row in files:
        parsed = json.loads(row.get("parse_data") or "{}")
        for reference in parsed.get("references", []):
            normalized = (Path(row["path"]).parent / reference).as_posix()
            if normalized in by_path:
                edges.append(
                    {
                        "source": row["id"],
                        "target": by_path[normalized]["id"],
                        "type": "reference",
                        "evidence": [f"{row['path']} references {reference}"],
                    }
                )
        for imported in parsed.get("imports", []):
            stem = imported.split(".")[-1].split("/")[-1].lower()
            for target_id in by_stem.get(stem, set()):
                if target_id != row["id"]:
                    edges.append(
                        {
                            "source": row["id"],
                            "target": target_id,
                            "type": "import",
                            "evidence": [f"{row['path']} imports {imported}"],
                        }
                    )
    return edges
