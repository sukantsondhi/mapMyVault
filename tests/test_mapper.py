import json
import io
import shutil
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch

from src.config import MapperConfig
from src.analyzer import _resolve_local_yolo_model, extract_content
from src.exporter import export_obsidian_from_index
from src.mapper import RepositoryMapper
from src.models import FileSummary, RelationshipDecision
from src.query import VaultIndex
from src.storage import IndexStore


class FakeAI:
    def __init__(self):
        self.summary_calls = 0
        self.embedding_calls = 0
        self.relationship_calls = 0

    def models(self):
        return ["llama3.1:8b", "nomic-embed-text:latest"]

    def summarize(self, path, content, parse_data, document_metadata="{}"):
        self.summary_calls += 1
        concepts = ["authentication"] if "token" in content else ["storage"]
        return FileSummary(
            title=Path(path).name,
            purpose=content[:80],
            detailed_description=f"Detailed explanation of {Path(path).name}",
            document_type=json.loads(document_metadata).get("format", "text"),
            responsibilities=["local analysis"],
            concepts=concepts,
            topics=concepts,
            tags=concepts,
            key_entities=["mapMyVault"],
            important_dates=[],
            interfaces=[],
            dependencies=json.loads(parse_data).get("imports", []),
            suggested_actions=["Review this file"],
        )

    def embed(self, texts):
        self.embedding_calls += len(texts)
        return [[1.0, 0.0, 0.0] for _ in texts]

    def judge_relationship(self, source_path, source_summary, target_path, target_summary):
        self.relationship_calls += 1
        return RelationshipDecision(
            related=True,
            relationship_type="content",
            confidence=0.9,
            explanation="Both files implement the same content-level responsibility.",
            evidence=["Both summaries describe local analysis."],
        )


class FakeVectorIndex:
    def __init__(self):
        self.ids = []

    def upsert(self, file_id, document, embedding, path):
        if file_id not in self.ids:
            self.ids.append(file_id)

    def candidates(self, embedding, limit):
        return [
            {"id": file_id, "distance": 0.1, "metadata": {}}
            for file_id in self.ids[:limit]
        ]

    def delete(self, file_ids):
        self.ids = [item for item in self.ids if item not in file_ids]

    def close(self):
        pass


class RepositoryMapperTests(unittest.TestCase):
    def test_mixed_pdf_preserves_native_and_scanned_text(self):
        import fitz

        with tempfile.TemporaryDirectory() as base:
            path = Path(base) / "mixed.pdf"
            with fitz.open() as scan:
                scan.new_page().insert_text((72, 100), "SCANNED REFERENCE 4286", fontsize=24)
                image = scan[0].get_pixmap().tobytes("png")
            with fitz.open() as document:
                document.new_page().insert_text((72, 100), "NATIVE REFERENCE 7391", fontsize=24)
                page = document.new_page()
                page.insert_image(page.rect, stream=image)
                document.save(path)
            with patch("src.analyzer._ocr_pdf_tesseract", return_value=(
                "SCANNED REFERENCE 4286", {"extraction_status": "ocr_extracted", "ocr_pages": 1}
            )) as ocr:
                text, metadata = extract_content(path, 10_000_000, 24_000, enable_ocr=True)
            self.assertIn("NATIVE REFERENCE 7391", text)
            self.assertIn("SCANNED REFERENCE 4286", text)
            self.assertEqual(metadata["page_count"], 2)
            ocr.assert_called_once()

    def test_local_office_document_extractors_read_content_and_metadata(self):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as base:
            base = Path(base)

            docx_path = base / "project_plan.docx"
            document = Document()
            document.core_properties.title = "Project Plan"
            document.add_paragraph("Launch the local knowledge database in June.")
            document.save(docx_path)

            xlsx_path = base / "budget.xlsx"
            workbook = Workbook()
            workbook.active.title = "Budget"
            workbook.active.append(["Item", "Cost"])
            workbook.active.append(["Local server", 500])
            workbook.save(xlsx_path)

            pptx_path = base / "briefing.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Local Search"
            slide.placeholders[1].text = "Query PDFs and documents without cloud services."
            presentation.save(pptx_path)

            docx_text, docx_meta = extract_content(docx_path, 10_000_000, 24_000)
            xlsx_text, xlsx_meta = extract_content(xlsx_path, 10_000_000, 24_000)
            pptx_text, pptx_meta = extract_content(pptx_path, 10_000_000, 24_000)

            self.assertIn("local knowledge database", docx_text)
            self.assertEqual(docx_meta["title"], "Project Plan")
            self.assertIn("Local server", xlsx_text)
            self.assertEqual(xlsx_meta["sheet_names"], ["Budget"])
            self.assertIn("without cloud services", pptx_text)
            self.assertEqual(pptx_meta["slide_count"], 1)

    def test_optional_ocr_reports_engine_errors_without_remote_fallback(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            image = base / "scan.jpg"
            image.write_bytes(b"fake-local-image")

            text, metadata = extract_content(
                image,
                10_000_000,
                24_000,
                enable_ocr=True,
                ocr_engine="cloud-ocr",
            )

            self.assertEqual(text, "")
            self.assertEqual(metadata["extraction_status"], "ocr_failed")
            self.assertIn("Unsupported OCR engine", metadata["ocr_error"])

    def test_yolo_vision_labels_are_extracted_as_searchable_text(self):
        class FakeArray:
            def tolist(self):
                return [1, 2, 30, 40]

        class FakeBox:
            cls = [0]
            conf = [0.91]
            xyxy = [FakeArray()]

        class FakeResult:
            names = {0: "person"}
            boxes = [FakeBox()]

        class FakeYOLO:
            names = {0: "person"}

            def __init__(self, model):
                self.model = model

            def predict(self, **kwargs):
                return [FakeResult()]

        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            image = base / "photo.jpg"
            image.write_bytes(b"fake-local-image")
            weights = base / "yolov8n.pt"
            weights.write_bytes(b"local-weights")
            fake_module = types.SimpleNamespace(YOLO=FakeYOLO)

            with patch.dict(sys.modules, {"ultralytics": fake_module}):
                text, metadata = extract_content(
                    image,
                    10_000_000,
                    24_000,
                    enable_vision=True,
                    vision_model=str(weights),
                )

            self.assertIn("Detected objects: person x1", text)
            self.assertEqual(metadata["extraction_status"], "vision_extracted")
            self.assertEqual(metadata["vision_status"], "vision_extracted")
            self.assertEqual(metadata["vision_label_counts"], {"person": 1})

    def test_yolo_empty_result_is_recorded_as_local_image_analysis(self):
        class FakeResult:
            names = {0: "person"}
            boxes = []

        class FakeYOLO:
            names = {0: "person"}

            def __init__(self, model):
                self.model = model

            def predict(self, **kwargs):
                return [FakeResult()]

        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            image = base / "alloy.jpg"
            image.write_bytes(b"fake-local-image")
            weights = base / "yolov8n.pt"
            weights.write_bytes(b"local-weights")
            fake_module = types.SimpleNamespace(YOLO=FakeYOLO)

            with patch.dict(sys.modules, {"ultralytics": fake_module}):
                text, metadata = extract_content(
                    image,
                    10_000_000,
                    24_000,
                    enable_vision=True,
                    vision_model=str(weights),
                )

            self.assertIn("No objects from the configured YOLO model", text)
            self.assertEqual(metadata["extraction_status"], "image_extracted")
            self.assertEqual(metadata["vision_status"], "vision_empty")

    def test_yolo_vision_requires_local_weights_without_auto_download(self):
        with tempfile.TemporaryDirectory() as base:
            image = Path(base) / "photo.jpg"
            image.write_bytes(b"fake-local-image")

            text, metadata = extract_content(
                image,
                10_000_000,
                24_000,
                enable_vision=True,
                vision_model="missing-yolo-model.pt",
            )

            self.assertEqual(text, "")
            self.assertEqual(metadata["vision_status"], "vision_unavailable")
            self.assertIn("not auto-download", metadata["vision_error"])

    def test_yolo_model_filename_resolves_from_repo_models_folder(self):
        with tempfile.TemporaryDirectory() as base:
            models = Path(base) / "models"
            models.mkdir()
            weights = models / "yolov8n.pt"
            weights.write_bytes(b"local-weights")

            with patch("src.analyzer.REPO_MODELS_DIR", models):
                resolved = _resolve_local_yolo_model("yolov8n.pt")

            self.assertEqual(resolved, weights.resolve())

    def test_mapper_stops_early_when_yolo_package_is_missing(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            (repo / "photo.jpg").write_bytes(b"fake-local-image")
            weights = base / "yolov8n.pt"
            weights.write_bytes(b"local-weights")

            mapper = RepositoryMapper(
                repo,
                output,
                MapperConfig(enable_vision=True, vision_model=str(weights)),
                ai=FakeAI(),
                use_chroma=False,
            )
            try:
                with patch("src.mapper.importlib.util.find_spec", return_value=None):
                    with self.assertRaises(RuntimeError) as error:
                        mapper.map(export_vault_notes=False)
            finally:
                mapper.close()

            self.assertIn("ultralytics", str(error.exception))
            self.assertIn("requirements-vision.txt", str(error.exception))

    def test_ocr_config_change_invalidates_extraction(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            source = repo / "scan.jpg"
            source.write_bytes(b"fake-local-image")
            ai = FakeAI()

            mapper = RepositoryMapper(repo, output, ai=ai, use_chroma=False)
            mapper.map()
            file_id = mapper.store.get_file_by_path("scan.jpg")["id"]
            first_metadata = json.loads(
                mapper.store.get_file(file_id)["document_metadata_json"]
            )
            mapper.close()

            mapper = RepositoryMapper(
                repo,
                output,
                config=MapperConfig(enable_ocr=True, ocr_engine="cloud-ocr"),
                ai=ai,
                use_chroma=False,
            )
            mapper.map()
            file_id = mapper.store.get_file_by_path("scan.jpg")["id"]
            second_metadata = json.loads(
                mapper.store.get_file(file_id)["document_metadata_json"]
            )
            mapper.close()

            self.assertEqual(first_metadata["extraction_status"], "ocr_required")
            self.assertEqual(second_metadata["extraction_status"], "ocr_failed")

    def test_mapper_warns_when_pdf_needs_ocr_but_ocr_is_disabled(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            (repo / "blank.pdf").write_bytes(
                b"%PDF-1.4\n"
                b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
                b"2 0 obj<</Type/Pages/Count 1/Kids[3 0 R]>>endobj\n"
                b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 72 72]>>endobj\n"
                b"trailer<</Root 1 0 R>>\n%%EOF"
            )
            warnings = []

            mapper = RepositoryMapper(
                repo,
                output,
                ai=FakeAI(),
                use_chroma=False,
                warning_callback=warnings.append,
            )
            try:
                mapper.map(export_vault_notes=False)
            finally:
                mapper.close()

            self.assertTrue(any("Extraction/OCR could not run" in warning for warning in warnings))

    def test_mapper_skips_file_level_ocr_failures(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            for index in range(3):
                (repo / f"scan-{index}.jpg").write_bytes(b"fake-local-image")

            config = MapperConfig(
                enable_ocr=True,
                ocr_engine="cloud-ocr",
            )
            warnings = []
            logs = []
            mapper = RepositoryMapper(
                repo,
                output,
                config,
                ai=FakeAI(),
                use_chroma=False,
                warning_callback=warnings.append,
                log_callback=logs.append,
            )
            try:
                mapper.map(export_vault_notes=False)
            finally:
                mapper.close()

            self.assertEqual(len(warnings), 3)
            self.assertTrue(any("indexing skipped" in line for line in logs))

    def test_mapper_stops_when_ocr_tool_is_unavailable(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            (repo / "scan.jpg").write_bytes(b"fake-local-image")

            mapper = RepositoryMapper(
                repo,
                output,
                MapperConfig(enable_ocr=True),
                ai=FakeAI(),
                use_chroma=False,
            )
            try:
                with patch("src.analyzer.shutil.which", return_value=None):
                    with self.assertRaises(RuntimeError) as error:
                        mapper.map(export_vault_notes=False)
            finally:
                mapper.close()

            self.assertIn("OCR is unavailable", str(error.exception))

    def test_mapper_logs_per_file_index_success(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            (repo / "notes.txt").write_text("local token notes", encoding="utf-8")
            logs = []

            mapper = RepositoryMapper(
                repo,
                output,
                ai=FakeAI(),
                use_chroma=False,
                log_callback=logs.append,
            )
            try:
                mapper.map(export_vault_notes=False)
            finally:
                mapper.close()

            self.assertTrue(any("text extracted" in line and "notes.txt" in line for line in logs))
            self.assertTrue(any("indexed successfully" in line and "notes.txt" in line for line in logs))

    def test_content_pipeline_outputs_vault_database_and_evidence_links(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            (repo / "alpha.py").write_text(
                "import beta\n\ndef issue_token():\n    return 'token'\n", encoding="utf-8"
            )
            (repo / "beta.py").write_text(
                "def validate_token(value):\n    return bool(value)\n", encoding="utf-8"
            )
            (repo / "docs").mkdir()
            (repo / "docs" / "guide.md").write_text(
                "# Guide\n\nThis explains token storage.", encoding="utf-8"
            )
            ai = FakeAI()
            mapper = RepositoryMapper(
                repo,
                output,
                MapperConfig(),
                ai=ai,
                vector_index=FakeVectorIndex(),
            )
            try:
                status = mapper.map()
            finally:
                mapper.close()

            self.assertEqual(status["files"], 3)
            self.assertTrue((output / "data" / "index.sqlite").exists())
            self.assertTrue((output / "data" / "manifest.json").exists())
            self.assertTrue((output / "obsidian" / "alpha.py.md").exists())
            self.assertTrue((output / "obsidian" / "docs" / "docs-index.md").exists())
            note = (output / "obsidian" / "alpha.py.md").read_text(encoding="utf-8")
            self.assertIn("alpha.py imports beta", note)
            self.assertIn("content-level responsibility", note)
            self.assertIn('"tags"', note)
            self.assertIn("#topic/authentication", note)
            folder_note = (output / "obsidian" / "docs" / "docs-index.md").read_text(
                encoding="utf-8"
            )
            self.assertIn("[[docs/guide.md]]", folder_note)

    def test_index_only_run_can_export_obsidian_later(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            (repo / "notes.txt").write_text("John payslip archive", encoding="utf-8")

            mapper = RepositoryMapper(repo, output, ai=FakeAI(), use_chroma=False)
            try:
                mapper.map(export_vault_notes=False)
            finally:
                mapper.close()

            self.assertTrue((output / "data" / "index.sqlite").exists())
            self.assertTrue((output / "data" / "graph.json").exists())
            self.assertFalse((output / "obsidian" / "INDEX.md").exists())

            result = export_obsidian_from_index(output)
            self.assertEqual(result["files"], 1)
            self.assertTrue((output / "obsidian" / "INDEX.md").exists())

    def test_index_subtree_can_be_soft_deleted(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            (repo / "keep.txt").write_text("keep token", encoding="utf-8")
            (repo / "old").mkdir()
            (repo / "old" / "remove.txt").write_text("remove token", encoding="utf-8")

            mapper = RepositoryMapper(repo, output, ai=FakeAI(), use_chroma=False)
            try:
                mapper.map(export_vault_notes=False)
            finally:
                mapper.close()

            store = IndexStore(output / "data" / "index.sqlite")
            try:
                removed = store.mark_path_tree_deleted("old")
                store.rebuild_fts()
                self.assertEqual(len(removed), 2)
                self.assertIsNone(store.get_file_by_path("old"))
                self.assertIsNone(store.get_file_by_path("old/remove.txt"))
                self.assertIsNotNone(store.get_file_by_path("keep.txt"))
                self.assertEqual(store.search("remove", 10), [])
            finally:
                store.close()

    def test_index_subtree_removal_treats_wildcards_literally(self):
        with tempfile.TemporaryDirectory() as base:
            repo = Path(base) / "repo"
            output = Path(base) / "output"
            for folder in ("old_%", "old_Xstuff"):
                (repo / folder).mkdir(parents=True)
                (repo / folder / "file.txt").write_text(folder, encoding="utf-8")
            mapper = RepositoryMapper(repo, output, ai=FakeAI(), use_chroma=False)
            try:
                mapper.map(export_vault_notes=False)
            finally:
                mapper.close()
            store = IndexStore(output / "data" / "index.sqlite")
            try:
                removed = store.mark_path_tree_deleted("old_%")
                self.assertEqual(len(removed), 2)
                self.assertIsNone(store.get_file_by_path("old_%/file.txt"))
                self.assertIsNotNone(store.get_file_by_path("old_Xstuff/file.txt"))
            finally:
                store.close()

    def test_file_progress_callback_reports_current_paths(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            (repo / "first.txt").write_text("first local file", encoding="utf-8")
            (repo / "second.txt").write_text("second local file", encoding="utf-8")
            seen = []

            def file_progress(label, path, done, total):
                seen.append((label, path, done, total))

            mapper = RepositoryMapper(
                repo,
                output,
                ai=FakeAI(),
                use_chroma=False,
                file_progress_callback=file_progress,
            )
            try:
                mapper.map(export_vault_notes=False)
            finally:
                mapper.close()

            paths = {item[1] for item in seen}
            self.assertIn("first.txt", paths)
            self.assertIn("second.txt", paths)
            self.assertTrue(all(item[2] <= item[3] for item in seen))

    def test_stage_progress_reports_start_and_finish(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            (repo / "first.txt").write_text("first local file", encoding="utf-8")
            seen = []

            def stage_progress(label, done, total):
                seen.append((label, done, total))

            mapper = RepositoryMapper(
                repo,
                output,
                ai=FakeAI(),
                use_chroma=False,
                progress_callback=stage_progress,
            )
            try:
                mapper.map(export_vault_notes=False)
            finally:
                mapper.close()

            labels = [item[0] for item in seen]
            self.assertIn("Running: Discovery", labels)
            self.assertIn("Finished: Manifest", labels)
            self.assertEqual(seen[-1][1], seen[-1][2])

    def test_count_questions_use_normalized_terms_and_undergrad_folder(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            undergrad = repo / "Offer Letters" / "Undergrad"
            postgrad = repo / "Offer Letters" / "Postgrad"
            undergrad.mkdir(parents=True)
            postgrad.mkdir(parents=True)
            (undergrad / "Brunel University.txt").write_text(
                "Sukant received a university offer for a BSc degree.",
                encoding="utf-8",
            )
            (undergrad / "City University.txt").write_text(
                "Sukant received an undergraduate university offer.",
                encoding="utf-8",
            )
            (postgrad / "Aberdeen University.txt").write_text(
                "Sukant received a postgraduate offer that mentions an undergraduate degree.",
                encoding="utf-8",
            )

            mapper = RepositoryMapper(repo, output, ai=FakeAI(), use_chroma=False)
            try:
                mapper.map(export_vault_notes=False)
            finally:
                mapper.close()

            index = VaultIndex(output)
            try:
                result = index.ask_mapmyvault(
                    "how many uni offers does sukant has for undergraduate degrees"
                )
                with patch.object(index, "search_files") as search:
                    empty_result = index.ask_mapmyvault("how many zzznonexistent?")
                search.assert_not_called()
                self.assertEqual(empty_result["found_count"], 0)
                self.assertEqual(empty_result["paths"], [])
                self.assertIn("0 matching files", empty_result["answer"])
            finally:
                index.close()

            self.assertEqual(result["found_count"], 2)
            self.assertEqual(len(result["paths"]), 2)

    def test_config_excludes_skip_selected_source_items(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            (repo / "keep.txt").write_text("keep token", encoding="utf-8")
            (repo / "skip.txt").write_text("skip token", encoding="utf-8")
            (repo / "Skip Folder").mkdir()
            (repo / "Skip Folder" / "nested.txt").write_text(
                "nested skip token", encoding="utf-8"
            )

            config = MapperConfig(
                excludes=[
                    "skip.txt",
                    "Skip Folder/",
                    "Skip Folder/**",
                ]
            )
            mapper = RepositoryMapper(repo, output, config, ai=FakeAI(), use_chroma=False)
            try:
                mapper.map(export_vault_notes=False)
                self.assertIsNotNone(mapper.store.get_file_by_path("keep.txt"))
                self.assertIsNone(mapper.store.get_file_by_path("skip.txt"))
                self.assertIsNone(mapper.store.get_file_by_path("Skip Folder"))
                self.assertIsNone(mapper.store.get_file_by_path("Skip Folder/nested.txt"))
            finally:
                mapper.close()

    def test_second_run_reuses_completed_stages_and_changed_file_is_partial(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            source = repo / "service.py"
            source.write_text("def first(): return 'token'\n", encoding="utf-8")
            ai = FakeAI()

            mapper = RepositoryMapper(repo, output, ai=ai, use_chroma=False)
            mapper.map()
            mapper.close()
            first_summary_calls = ai.summary_calls

            mapper = RepositoryMapper(repo, output, ai=ai, use_chroma=False)
            mapper.map()
            mapper.close()
            self.assertEqual(ai.summary_calls, first_summary_calls)

            source.write_text("def second(): return 'token changed'\n", encoding="utf-8")
            mapper = RepositoryMapper(repo, output, ai=ai, use_chroma=False)
            mapper.map()
            mapper.close()
            self.assertEqual(ai.summary_calls, first_summary_calls + 1)

            index = VaultIndex(output)
            self.assertIn("second", index.read_file_excerpt("service.py"))
            index.close()

    def test_output_inside_repository_is_rejected(self):
        with tempfile.TemporaryDirectory() as base:
            repo = Path(base)
            with self.assertRaises(ValueError):
                RepositoryMapper(repo, repo / "output", ai=FakeAI(), use_chroma=False)

    def test_missing_cached_artifact_is_repaired_and_deleted_note_is_removed(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            source = repo / "repair.txt"
            source.write_text("token content", encoding="utf-8")
            ai = FakeAI()
            mapper = RepositoryMapper(repo, output, ai=ai, use_chroma=False)
            mapper.map()
            file_id = mapper.store.get_file_by_path("repair.txt")["id"]
            mapper.store.connection.execute(
                "UPDATE files SET summary_json=NULL WHERE id=?", (file_id,)
            )
            mapper.store.connection.commit()
            mapper.close()

            calls = ai.summary_calls
            mapper = RepositoryMapper(repo, output, ai=ai, use_chroma=False)
            mapper.map()
            mapper.close()
            self.assertEqual(ai.summary_calls, calls + 1)

            source.unlink()
            mapper = RepositoryMapper(repo, output, ai=ai, use_chroma=False)
            mapper.map()
            mapper.close()
            self.assertFalse((output / "obsidian" / "repair.txt.md").exists())

    def test_document_details_are_persisted_searchable_and_exported(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            repo = base / "repo"
            output = base / "output"
            repo.mkdir()
            document = Document()
            document.core_properties.author = "Local Author"
            document.add_paragraph("The Phoenix project stores customer invoices locally.")
            document.save(repo / "phoenix.docx")

            mapper = RepositoryMapper(repo, output, ai=FakeAI(), use_chroma=False)
            mapper.map()
            mapper.close()

            index = VaultIndex(output)
            details = index.get_file_summary("phoenix.docx")
            formats = index.list_files_by_format("docx")
            answer = index.answer_question("What documents mention Phoenix invoices?", limit=5)
            count = index.count_matches("How many Phoenix invoices?", limit=10)
            direct = index.ask_mapmyvault("How many Phoenix invoices?", limit=10)
            index.close()

            self.assertEqual(details["document_metadata"]["author"], "Local Author")
            self.assertEqual(details["summary"]["document_type"], "docx")
            self.assertEqual(formats[0]["path"], "phoenix.docx")
            self.assertGreaterEqual(answer["found_count"], 1)
            self.assertEqual(answer["results"][0]["path"], "phoenix.docx")
            self.assertIn("Phoenix", answer["results"][0]["evidence_excerpt"])
            self.assertEqual(count["count"], 1)
            self.assertEqual(count["paths"], ["phoenix.docx"])
            self.assertEqual(direct["found_count"], 1)
            self.assertIn("contains 1 matching file", direct["answer"])
            self.assertEqual(direct["paths"], ["phoenix.docx"])
            note = (output / "obsidian" / "phoenix.docx.md").read_text(encoding="utf-8")
            self.assertIn("Detailed Description", note)
            self.assertIn("Local Author", note)
            self.assertIn("#type/docx", note)

    def test_blank_ocr_language_defaults_to_eng(self):
        with tempfile.TemporaryDirectory() as base:
            base = Path(base)
            source = base / "scan.jpg"
            source.write_bytes(b"not-a-real-image")

            text, metadata = extract_content(
                source,
                10_000_000,
                24_000,
                enable_ocr=True,
                ocr_language="",
            )

            self.assertEqual(text, "")
            self.assertEqual(metadata["ocr_language"], "eng")


class TextRecognitionFixture(unittest.TestCase):
    native_marker = "NATIVE REFERENCE 7391"
    scanned_marker = "SCANNED REFERENCE 4286"

    def setUp(self):
        import fitz

        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.base = Path(temporary.name)
        with fitz.open() as document:
            page = document.new_page(width=600, height=200)
            page.insert_text((30, 65), self.scanned_marker, fontsize=26)
            page.insert_text((30, 115), "Invoice total 1250", fontsize=24)
            self.scan_bytes = page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes("png")

    def _make_pdf(self, name, layouts):
        import fitz

        path = self.base / name
        with fitz.open() as document:
            for native, scanned in layouts:
                page = document.new_page(width=612, height=792)
                if native:
                    page.insert_text((36, 65), self.native_marker, fontsize=22)
                if scanned:
                    page.insert_image(fitz.Rect(36, 120, 576, 300), stream=self.scan_bytes)
            document.set_metadata({"title": "Recognition fixture", "author": "Test suite"})
            document.save(path)
        return path

    def _extract(self, path, **options):
        return extract_content(path, 10_000_000, options.pop("max_chars", 24_000), **options)

    def _assert_recognized(self, text, marker):
        self.assertIn(marker, " ".join(text.upper().split()))


class TextRecognitionTests(TextRecognitionFixture):
    def test_text_files_and_native_pdf_do_not_call_ocr(self):
        plain = self.base / "plain.txt"
        plain.write_text("PLAINTEXT REFERENCE 3815", encoding="utf-8-sig")
        native = self._make_pdf("native.pdf", [(True, False)])
        with patch("pytesseract.image_to_string", side_effect=AssertionError("OCR was not needed")):
            for path, marker in ((plain, "PLAINTEXT REFERENCE 3815"), (native, self.native_marker)):
                with self.subTest(path=path.name):
                    text, metadata = self._extract(path, enable_ocr=True)
                    self._assert_recognized(text, marker)
                    self.assertEqual(metadata["extraction_status"], "extracted")
                    self.assertEqual(metadata["extracted_characters"], len(text))

    def test_disabled_ocr_retains_native_text_and_reports_scanned_pages(self):
        path = self._make_pdf("mixed.pdf", [(True, False), (False, True)])
        with patch("pytesseract.image_to_string", side_effect=AssertionError("OCR is disabled")):
            text, metadata = self._extract(path)
        self._assert_recognized(text, self.native_marker)
        self.assertNotIn(self.scanned_marker, text)
        self.assertEqual(metadata["ocr_status"], "ocr_required")
        self.assertEqual(metadata["ocr_candidate_pages"], [2])

    def test_missing_engine_keeps_native_text(self):
        path = self._make_pdf("mixed.pdf", [(True, True)])
        with patch("src.analyzer.shutil.which", return_value=None):
            text, metadata = self._extract(path, enable_ocr=True)
        self._assert_recognized(text, self.native_marker)
        self.assertEqual(metadata["ocr_status"], "ocr_unavailable")

    def test_disabled_image_ocr_is_reported_as_required(self):
        path = self.base / "image.png"
        path.write_bytes(self.scan_bytes)
        text, metadata = self._extract(path)
        self.assertEqual(text, "")
        self.assertEqual(metadata["extraction_status"], "ocr_required")

    def test_native_pdf_character_limit_stops_page_extraction(self):
        path = self._make_pdf("native-limit.pdf", [(True, False), (True, False)])
        with patch("pypdf._page.PageObject.extract_text", side_effect=[
            self.native_marker, AssertionError("Character limit already reached")
        ]) as native:
            text, metadata = self._extract(path, max_chars=10, enable_ocr=True)
        native.assert_called_once()
        self.assertEqual(len(text), 10)
        self.assertTrue(metadata["truncated"])

    def test_ocr_page_limit_can_exclude_all_scanned_pages(self):
        path = self._make_pdf("page-limit.pdf", [(True, False), (False, True)])
        with patch("src.analyzer._ocr_pdf_tesseract", side_effect=AssertionError("Page outside OCR limit")):
            text, metadata = self._extract(path, enable_ocr=True, ocr_max_pages=1)
        self._assert_recognized(text, self.native_marker)
        self.assertEqual(metadata["ocr_pages"], 0)
        self.assertEqual(metadata["ocr_skipped_pages"], [2])
        self.assertEqual(metadata["ocr_status"], "ocr_skipped")

    def test_same_page_ocr_does_not_duplicate_native_lines(self):
        path = self._make_pdf("same-page.pdf", [(True, True)])
        with patch("src.analyzer._ocr_pdf_tesseract", return_value=("", {
            "extraction_status": "ocr_extracted",
            "ocr_page_texts": {1: "native   reference 7391\n" + self.scanned_marker},
        })):
            text, metadata = self._extract(path, enable_ocr=True)
        self.assertEqual(text, self.native_marker + "\n" + self.scanned_marker)
        self.assertNotIn("ocr_page_texts", metadata)

    def test_old_ocr_version_refreshes_cached_mixed_pdf(self):
        repository = self.base / "source"
        repository.mkdir()
        self._make_pdf("source/mixed.pdf", [(True, False), (False, True)])
        config = MapperConfig(enable_ocr=True)
        ai = FakeAI()
        mapper = RepositoryMapper(repository, self.base / "output", config, ai=ai, use_chroma=False)
        try:
            with patch("src.mapper.extract_content", return_value=(
                self.native_marker, {"format": "pdf", "extraction_status": "extracted"}
            )):
                mapper.map(export_vault_notes=False)
            first_calls = ai.summary_calls
            mapper.store.set_metadata("ocr_version", (
                f"{config.ocr_engine}:{config.ocr_language}:"
                f"{config.ocr_dpi}:{config.ocr_max_pages}:{config.enable_ocr}"
            ))
            mapper.store.connection.commit()
            with patch("src.analyzer._ocr_pdf_tesseract", return_value=(self.scanned_marker, {
                "extraction_status": "ocr_extracted",
                "ocr_pages": 1,
                "ocr_page_texts": {2: self.scanned_marker},
            })) as ocr:
                mapper.map(export_vault_notes=False)
            ocr.assert_called_once()
            self.assertGreater(ai.summary_calls, first_calls)
            text = mapper.store.get_file_by_path("mixed.pdf")["extracted_text"]
            self._assert_recognized(text, self.native_marker)
            self._assert_recognized(text, self.scanned_marker)
        finally:
            mapper.close()


@unittest.skipUnless(shutil.which("tesseract"), "Live OCR requires Tesseract on PATH with English language data")
class LiveTextRecognitionTests(TextRecognitionFixture):
    def test_real_image_ocr_for_supported_formats(self):
        from PIL import Image

        for suffix in ("png", "jpg", "tiff", "bmp", "webp"):
            with self.subTest(format=suffix):
                path = self.base / f"image.{suffix}"
                with Image.open(io.BytesIO(self.scan_bytes)) as image:
                    image.convert("RGB").save(path)
                text, metadata = self._extract(path, enable_ocr=True)
                self._assert_recognized(text, self.scanned_marker)
                self._assert_recognized(text, "INVOICE TOTAL 1250")
                self.assertEqual(metadata["ocr_status"], "ocr_extracted")
                self.assertEqual(metadata["extracted_characters"], len(text))

    def test_real_scanned_pdf_ocr(self):
        path = self._make_pdf("scan.pdf", [(False, True)])
        text, metadata = self._extract(path, enable_ocr=True)
        self._assert_recognized(text, self.scanned_marker)
        self.assertEqual(metadata["extraction_status"], "ocr_extracted")
        self.assertEqual(metadata["ocr_pages"], 1)
        self.assertEqual(metadata["title"], "Recognition fixture")

    def test_real_mixed_pdf_pages_preserve_order(self):
        for layouts in ([(True, False), (False, True)], [(False, True), (True, False)]):
            with self.subTest(layouts=layouts):
                path = self._make_pdf("mixed.pdf", layouts)
                text, metadata = self._extract(path, enable_ocr=True)
                self._assert_recognized(text, self.native_marker)
                self._assert_recognized(text, self.scanned_marker)
                first, second = (self.native_marker, self.scanned_marker) if layouts[0][0] else (self.scanned_marker, self.native_marker)
                self.assertLess(text.index(first), text.index(second))
                self.assertEqual(metadata["ocr_pages"], 1)

    def test_real_same_page_native_text_and_image_text(self):
        path = self._make_pdf("same-page.pdf", [(True, True)])
        text, metadata = self._extract(path, enable_ocr=True)
        self._assert_recognized(text, self.native_marker)
        self._assert_recognized(text, self.scanned_marker)
        self.assertEqual(text.count(self.native_marker), 1)
        self.assertEqual(metadata["ocr_pages"], 1)

    def test_real_pdf_renderer_fallback(self):
        path = self._make_pdf("fallback.pdf", [(False, True)])
        with patch("src.analyzer._pdf_images_with_pdf2image", side_effect=RuntimeError("No Poppler")):
            text, metadata = self._extract(path, enable_ocr=True)
        self._assert_recognized(text, self.scanned_marker)
        self.assertEqual(metadata["ocr_pdf_renderer"], "pymupdf")

    def test_real_pdf_limits_are_reported(self):
        path = self._make_pdf("limits.pdf", [(False, True), (False, True), (True, False)])
        text, metadata = self._extract(path, enable_ocr=True, ocr_max_pages=1)
        self.assertEqual(metadata["ocr_pages"], 1)
        self.assertEqual(metadata["ocr_skipped_pages"], [2])
        self._assert_recognized(text, self.native_marker)
        text, metadata = self._extract(path, enable_ocr=True, max_chars=30)
        self.assertLessEqual(len(text), 30)
        self.assertTrue(metadata["truncated"])
        self.assertEqual(metadata["ocr_pages"], 1)

    def test_real_blank_and_invalid_images(self):
        from PIL import Image

        blank = self.base / "blank.png"
        with Image.new("RGB", (600, 200), "white") as image:
            image.save(blank)
        text, metadata = self._extract(blank, enable_ocr=True)
        self.assertEqual(text, "")
        self.assertEqual(metadata["ocr_status"], "ocr_empty")
        invalid = self.base / "invalid.png"
        invalid.write_bytes(b"not an image")
        text, metadata = self._extract(invalid, enable_ocr=True)
        self.assertEqual(text, "")
        self.assertEqual(metadata["ocr_status"], "ocr_failed")

    def test_real_missing_language_is_reported(self):
        path = self.base / "image.png"
        path.write_bytes(self.scan_bytes)
        text, metadata = self._extract(path, enable_ocr=True, ocr_language="nonexistent_language")
        self.assertEqual(text, "")
        self.assertEqual(metadata["ocr_status"], "ocr_failed")
        self.assertIn("nonexistent_language", metadata["ocr_error"])

    def test_real_ocr_text_is_persisted_searchable_and_reused(self):
        repository = self.base / "source"
        repository.mkdir()
        self._make_pdf("source/mixed.pdf", [(True, False), (False, True)])
        (repository / "scan.png").write_bytes(self.scan_bytes)
        (repository / "plain.txt").write_text("PLAINTEXT REFERENCE 3815", encoding="utf-8")
        output = self.base / "output"
        ai = FakeAI()
        mapper = RepositoryMapper(repository, output, MapperConfig(enable_ocr=True), ai=ai, use_chroma=False)
        try:
            mapper.map(export_vault_notes=False)
            first_calls = ai.summary_calls
            with patch("pytesseract.image_to_string", side_effect=AssertionError("Unchanged OCR should be cached")):
                mapper.map(export_vault_notes=False)
            self.assertEqual(ai.summary_calls, first_calls)
        finally:
            mapper.close()
        index = VaultIndex(output)
        try:
            for path in ("mixed.pdf", "scan.png"):
                self._assert_recognized(index.read_file_excerpt(path), self.scanned_marker)
            self._assert_recognized(index.read_file_excerpt("mixed.pdf"), self.native_marker)
            paths = {row["path"] for row in index.store.search("4286", 10)}
            self.assertEqual(paths, {"mixed.pdf", "scan.png"})
            self.assertIn("3815", index.read_file_excerpt("plain.txt"))
        finally:
            index.close()


if __name__ == "__main__":
    unittest.main()
